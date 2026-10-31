from pathlib import Path
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "docs" / "FINAL_PROJECT_REPORT.md"
OUT_DOCX = ROOT / "OmniSight_Final_Internship_Report.docx"
OUT_PDF = ROOT / "OmniSight_Final_Internship_Report.pdf"


def parse_markdown(text):
    lines = text.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        if line.startswith("```"):
            code, i = [], i + 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            blocks.append(("code", "\n".join(code)))
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[i + 1]):
            raw_rows = [line]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw_rows.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in row.strip("|").split("|")] for row in raw_rows]
            blocks.append(("table", rows))
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.*)", line)
        if bullet:
            items = []
            while i < len(lines):
                m = re.match(r"^\s*[-*]\s+(.*)", lines[i])
                if not m:
                    break
                items.append(m.group(1).strip())
                i += 1
            blocks.append(("bullets", items))
            continue

        numbered = re.match(r"^\s*\d+\.\s+(.*)", line)
        if numbered:
            items = []
            while i < len(lines):
                m = re.match(r"^\s*\d+\.\s+(.*)", lines[i])
                if not m:
                    break
                items.append(m.group(1).strip())
                i += 1
            blocks.append(("numbers", items))
            continue

        para = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if (not nxt.strip() or nxt.strip() == "---" or nxt.startswith("#")
                or nxt.startswith("```") or nxt.startswith("|")
                or re.match(r"^\s*[-*]\s+", nxt)
                or re.match(r"^\s*\d+\.\s+", nxt)):
                break
            para.append(nxt.strip())
            i += 1
        blocks.append(("paragraph", " ".join(para)))
    return blocks


def add_docx_inline(p, text):
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|\*.*?\*|\[[^\]]+\]\([^)]+\))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            p.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = p.add_run(token[1:-1])
            run.font.name, run.font.size = "Consolas", Pt(9)
        elif token.startswith("*"):
            p.add_run(token[1:-1]).italic = True
        else:
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            run = p.add_run(link_match.group(1))
            run.underline = True
            run.font.color.rgb = RGBColor(0, 102, 153)
        cursor = match.end()
    if cursor < len(text):
        p.add_run(text[cursor:])


def pdf_inline(text):
    value = escape(text)
    value = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"`(.+?)`", r'<font name="Courier">\1</font>', value)
    value = re.sub(r"\*(.+?)\*", r"<i>\1</i>", value)
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<link href="\2" color="#006699"><u>\1</u></link>', value)
    return value


def make_docx(blocks):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    for style_name, size, hex_color in [
        ("Title", 22, "16324F"),
        ("Heading 1", 16, "16324F"),
        ("Heading 2", 13, "245B78"),
        ("Heading 3", 11, "245B78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(hex_color)

    for kind, data in blocks:
        if kind == "heading":
            level, text = data
            if level == 1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_paragraph(style=f"Heading {min(level - 1, 3)}")
            add_docx_inline(p, text)
        elif kind == "paragraph":
            add_docx_inline(doc.add_paragraph(), data)
        elif kind in ("bullets", "numbers"):
            style = "List Bullet" if kind == "bullets" else "List Number"
            for item in data:
                add_docx_inline(doc.add_paragraph(style=style), item)
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(data)
            run.font.name, run.font.size = "Consolas", Pt(8)
        elif kind == "table":
            if not data:
                continue
            columns = max(len(row) for row in data)
            table = doc.add_table(rows=0, cols=columns)
            table.style = "Light Shading Accent 1"
            for row_index, row in enumerate(data):
                cells = table.add_row().cells
                for col_index in range(columns):
                    value = row[col_index] if col_index < len(row) else ""
                    cells[col_index].text = re.sub(r"<br\s*/?>", "\n", value, flags=re.I)
                    for p in cells[col_index].paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(8)
                            if row_index == 0:
                                run.bold = True

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("OmniSight - Final Internship Project Report")
    doc.save(OUT_DOCX)


def make_pdf(blocks):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=20, leading=25, alignment=TA_CENTER,
        textColor=colors.HexColor("#16324F"), spaceAfter=16
    ))
    styles.add(ParagraphStyle(
        name="ReportH1", parent=styles["Heading1"], fontName="Helvetica-Bold",
        fontSize=15, leading=19, textColor=colors.HexColor("#16324F"),
        spaceBefore=12, spaceAfter=7, keepWithNext=True
    ))
    styles.add(ParagraphStyle(
        name="ReportH2", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=11, leading=14, textColor=colors.HexColor("#245B78"),
        spaceBefore=8, spaceAfter=4, keepWithNext=True
    ))
    styles.add(ParagraphStyle(
        name="ReportBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9, leading=12, alignment=TA_LEFT, spaceAfter=6
    ))
    styles.add(ParagraphStyle(
        name="ReportBullet", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9, leading=12, leftIndent=14, firstLineIndent=-8, spaceAfter=3
    ))
    styles.add(ParagraphStyle(
        name="ReportCode", parent=styles["Code"], fontName="Courier",
        fontSize=7, leading=9, leftIndent=7, rightIndent=7, spaceAfter=6
    ))

    story = []
    for kind, data in blocks:
        if kind == "heading":
            level, text = data
            style = "ReportTitle" if level == 1 else ("ReportH1" if level <= 2 else "ReportH2")
            story.append(Paragraph(pdf_inline(text), styles[style]))
        elif kind == "paragraph":
            story.append(Paragraph(pdf_inline(data), styles["ReportBody"]))
        elif kind in ("bullets", "numbers"):
            for idx, item in enumerate(data, 1):
                prefix = "&#8226;&nbsp;" if kind == "bullets" else f"{idx}.&nbsp;"
                story.append(Paragraph(prefix + pdf_inline(item), styles["ReportBullet"]))
        elif kind == "code":
            story.append(Paragraph(escape(data).replace("\n", "<br/>"), styles["ReportCode"]))
        elif kind == "table":
            if not data:
                continue
            cols = max(len(row) for row in data)
            normalized = [row + [""] * (cols - len(row)) for row in data]
            rows = [[Paragraph(pdf_inline(cell), styles["ReportBody"]) for cell in row] for row in normalized]
            usable_width = A4[0] - 30 * mm
            table = Table(rows, colWidths=[usable_width / cols] * cols, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE8F0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AAAB5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([table, Spacer(1, 6)])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawCentredString(A4[0] / 2, 10 * mm,
                                 f"OmniSight - Final Internship Project Report | Page {doc.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(
        str(OUT_PDF), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm,
        topMargin=16 * mm, bottomMargin=18 * mm,
        title="OmniSight Final Internship Project Report", author="Sheefa J Pathan"
    )
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def main():
    if not SOURCE.exists():
        raise SystemExit(f"Source report not found: {SOURCE}")
    if OUT_DOCX.exists() or OUT_PDF.exists():
        raise SystemExit("An output already exists. Refusing to overwrite any existing final report.")
    source_text = SOURCE.read_text(encoding="utf-8-sig")
    blocks = parse_markdown(source_text)
    make_docx(blocks)
    make_pdf(blocks)
    print("Created:")
    print(f"  {OUT_DOCX}")
    print(f"  {OUT_PDF}")
    print(f"Parsed content blocks: {len(blocks)}")
    print("Please visually inspect both outputs before submission, especially tables and page breaks.")


if __name__ == "__main__":
    main()
